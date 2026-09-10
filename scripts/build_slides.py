"""Build the six-slide, business-facing executive decision narrative."""

import json
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "deliverables/london-office-market-agent.pptx"
NAVY = RGBColor(16, 35, 49)
CREAM = RGBColor(244, 241, 234)
TEAL = RGBColor(27, 128, 119)
MUTED = RGBColor(82, 103, 109)
WHITE = RGBColor(255, 253, 248)
PALE = RGBColor(173, 218, 209)
LINE = RGBColor(215, 213, 203)


def box(slide, x, y, w, h, text, size=20, color=NAVY, bold=False):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(0.02)
    frame.margin_top = frame.margin_bottom = 0
    frame.vertical_anchor = MSO_ANCHOR.TOP
    for i, line in enumerate(text.split("\n")):
        para = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        para.space_after = Pt(8)
        run = para.add_run()
        run.text = line
        run.font.name = "Arial"
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
    return shape


def panel(slide, x, y, w, h, dark=False):
    shape = slide.shapes.add_shape(1, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(25, 51, 64) if dark else WHITE
    shape.line.color.rgb = RGBColor(56, 84, 94) if dark else LINE


def base(prs, number, section, title, subtitle, dark=False):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = NAVY if dark else CREAM
    box(
        slide,
        0.65,
        0.36,
        10,
        0.22,
        f"LONDON MARKET MONITOR  /  {section.upper()}",
        10,
        PALE if dark else TEAL,
        True,
    )
    box(slide, 0.65, 0.96, 12, 1.15, title, 29, WHITE if dark else NAVY, True)
    box(
        slide, 0.67, 2.2 if number == 1 else 1.85, 11.9, 0.65, subtitle, 17, PALE if dark else MUTED
    )
    box(
        slide,
        0.67,
        7.02,
        10.5,
        0.22,
        "EXECUTIVE DECISION BRIEF  |  10 SEPTEMBER 2026",
        9,
        PALE if dark else MUTED,
    )
    box(slide, 12.0, 7.0, 0.7, 0.24, f"{number:02d} / 06", 10, PALE if dark else MUTED)
    return slide


def note(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    prs.core_properties.title = "London Market Monitor - executive decision brief"
    prs.core_properties.author = "London Market Monitor"

    s = base(
        prs,
        1,
        "problem",
        "Turn London office market noise into a\ncited, decision-ready brief",
        "The challenge is turning fragmented evidence into a defensible meeting position.",
        True,
    )
    for i, (title, body) in enumerate(
        [
            (
                "Fragmented evidence",
                "Broker reports, market news and economic signals arrive in different "
                "formats and periods.",
            ),
            (
                "Decision risk",
                "A headline rent or isolated deal can obscure the wider picture of "
                "demand, supply and vacancy.",
            ),
            (
                "A clearer position",
                "Bring the takeaway, supporting sources and unresolved questions into the "
                "same conversation.",
            ),
        ]
    ):
        x = 0.7 + i * 4.15
        panel(s, x, 2.85, 3.9, 2.55, True)
        box(s, x + 0.22, 3.08, 3.45, 0.4, title, 19, PALE, True)
        box(s, x + 0.22, 3.7, 3.42, 1.45, body, 18, WHITE)
    box(
        s,
        0.72,
        5.85,
        11.8,
        0.4,
        "Rents  /  Vacancy  /  Leasing demand  /  New supply  /  Economic drivers",
        19,
        WHITE,
    )
    box(
        s,
        0.72,
        6.4,
        11.8,
        0.3,
        "London focus: City, West End, Canary Wharf and Midtown / Fringe",
        16,
        PALE,
    )
    note(
        s,
        "Executive question: Why does this matter? The assessment calls for a runnable "
        "Python PoC and a 3-6 slide business deck. The value proposition is a traceable basis "
        "for discussion, without claiming measured time savings or investment advice. "
        "Vacancy means empty space; take-up means space leased. Coverage depends on "
        "accessible reports.",
    )

    s = base(
        prs,
        2,
        "decision",
        "Start with the decision you need to prepare for",
        "Define the business question, the market and the period before assessing the evidence.",
    )
    for i, (label, title, question, result) in enumerate(
        [
            (
                "MONITOR",
                "Reset priorities",
                "What changed versus the previous quarter?",
                "Identify material movements and decide what needs attention.",
            ),
            (
                "INVESTIGATE",
                "Challenge a view",
                "Is West End outperforming City?",
                "Test the evidence for and against a view before adopting it.",
            ),
            (
                "PREPARE",
                "Align the team",
                "What are the key developments this month?",
                "Agree the meeting position, open questions and what to watch.",
            ),
        ]
    ):
        x = 0.7 + i * 4.15
        panel(s, x, 2.75, 3.9, 3.8)
        box(s, x + 0.24, 3, 3.4, 0.25, label, 11, TEAL, True)
        box(s, x + 0.24, 3.43, 3.4, 0.7, title, 23, NAVY, True)
        box(s, x + 0.24, 4.33, 3.4, 0.85, question, 19)
        box(s, x + 0.24, 5.42, 3.4, 0.95, result, 17, MUTED)
    note(
        s,
        "Executive question: Which decision will this support? Preserve the three existing "
        "workflows, but frame each around its business use. Monitor needs a comparison period "
        "or a saved refresh baseline. Investigate tests a hypothesis and may return insufficient "
        "evidence. Prepare returns up to five dated developments when coverage supports them. "
        "Higher rent alone does not establish outperformance.",
    )

    s = base(
        prs,
        3,
        "how it works",
        "Bring your sources. Ask your question. Get cited answers.",
        "A question-led workflow turns accessible reports into a brief you can inspect.",
    )
    stages = [
        ("Sources", "Public reports, market news and supplied text"),
        ("Retrieve / Filter", "Find evidence relevant to the question and period"),
        ("Compare / Reason", "Match measures and weigh the findings"),
        ("Verify", "Check citations and numerical support"),
        ("Cited Brief", "State the takeaway, sources and gaps"),
    ]
    for i, (title, body) in enumerate(stages):
        x = 0.7 + i * 2.52
        panel(s, x, 3.05, 2.2, 2.25)
        box(s, x + 0.14, 3.3, 1.94, 0.55, title, 17, TEAL, True)
        box(s, x + 0.14, 4.05, 1.9, 1.05, body, 16)
        if i < 4:
            arrow = s.shapes.add_shape(
                33, Inches(x + 2.23), Inches(3.95), Inches(0.24), Inches(0.25)
            )
            arrow.fill.solid()
            arrow.fill.fore_color.rgb = TEAL
            arrow.line.fill.background()
    box(
        s,
        0.75,
        5.8,
        11.8,
        0.4,
        "Your team sets the question. The agent assembles evidence. People make the decision.",
        20,
        NAVY,
        True,
    )
    box(
        s,
        0.75,
        6.4,
        11.8,
        0.3,
        "Source passages stay available for review; gaps remain visible in the answer.",
        17,
        MUTED,
    )
    note(
        s,
        "Executive question: How does it work? This is a simplified logical workflow, "
        "not a literal sequence of one-time calls. The agent can retrieve, search and read "
        "again as needed within bounded calls. Supplied text can be ingested through the API; "
        "the PoC does not provide a general document-upload interface or private-data connector. "
        "Stored evidence uses semantic retrieval; validated numerical metrics use SQLite and "
        "deterministic calculations. Verification catches some citation and numerical errors; "
        "it does not guarantee source truth or every inference. Provider controls are omitted "
        "because they do not help the executive decision.",
    )

    s = base(
        prs,
        4,
        "output",
        "See the takeaway first. Go deeper only when you need to.",
        "Live example: do publishers disagree on City and West End prime rents?",
        True,
    )
    panel(s, 0.7, 2.7, 7.45, 3.25, True)
    box(s, 0.95, 2.95, 6.9, 0.25, "RECORDED AGENT OUTPUT / EDITED EXCERPT", 11, PALE, True)
    box(
        s,
        0.95,
        3.45,
        6.9,
        1.4,
        "Savills reports West End average prime rent of £175.10 per sq ft for Q2 2026. "
        "Knight Frank's £185.00 figure is a Q4 2025 West End Core headline rent. "
        "Different periods and definitions mean these are not a like-for-like "
        "disagreement. [1] [2]",
        21,
        WHITE,
    )
    box(
        s,
        0.95,
        5.35,
        6.85,
        0.4,
        "Run status: partial brief; no verified overall synthesis.",
        14,
        PALE,
        True,
    )
    box(s, 8.55, 2.95, 3.95, 0.4, "What this supports", 20, PALE, True)
    box(
        s,
        8.55,
        3.5,
        3.8,
        1.05,
        "Check definitions before treating different rent figures as conflicting evidence.",
        18,
        WHITE,
    )
    box(s, 8.55, 4.9, 3.95, 0.4, "What remains open", 20, PALE, True)
    box(
        s,
        8.55,
        5.45,
        3.8,
        1.05,
        "Knight Frank's Q2 rent table was absent from the retrieved text. The comparison "
        "remains open.",
        17,
        WHITE,
    )
    response = json.loads((ROOT / "deliverables/live-run.json").read_text())["response"]
    for y, text, index in [
        (6.18, "[1] Savills | Central London Office Market Watch, Q2 2026", 0),
        (6.52, "[2] Knight Frank | London Office Market Report, Q4 2025", 2),
    ]:
        shape = box(s, 0.95, y, 7.0, 0.25, text, 11, PALE)
        shape.click_action.hyperlink.address = response["citations"][index]["source"]["url"]
    note(
        s,
        "Executive question: What do I receive? This is an edited excerpt from the actual "
        "10 September 2026 run in deliverables/live-run.json, not a new model answer or an "
        "independently verified market assertion. Slide citation [1] maps to run citation [1]; "
        "slide [2] maps to run [3]. Links open the publishers. The run took about 105 seconds "
        "and returned a partial brief after a claim failed checks. Its missing overall synthesis "
        "is preserved visibly. The missing Knight Frank Q2 table was a retrieval limitation, "
        "not proof that the publisher did not report rents. Normally the opening brief connects "
        "findings, followed by detail and evidence. This example demonstrates a useful supported "
        "finding while keeping the unresolved comparison explicit.",
    )

    s = base(
        prs,
        5,
        "trust",
        "Check the claim before it reaches the meeting",
        "Trust comes from showing what can be compared, what conflicts and what remains unknown.",
    )
    for i, (title, body) in enumerate(
        [
            (
                "Comparable sources",
                "Match the measure, location, units and period before comparing values.",
            ),
            (
                "Conflicting evidence",
                "Keep competing reports visible. Do not average away differences in "
                "definitions or coverage.",
            ),
            (
                "Explicit uncertainty",
                "Separate reported facts from interpretation. Flag missing evidence and "
                "qualify the conclusion.",
            ),
            (
                "Traceable claims",
                "Follow citations to the publisher and inspect supporting passages, "
                "values and calculations.",
            ),
        ]
    ):
        x = 0.7 + (i % 2) * 6.3
        y = 2.75 + (i // 2) * 1.65
        panel(s, x, y, 5.95, 1.45)
        box(s, x + 0.22, y + 0.18, 5.5, 0.35, title, 20, TEAL, True)
        box(s, x + 0.22, y + 0.67, 5.45, 0.68, body, 17)
    box(
        s,
        0.75,
        6.35,
        11.85,
        0.4,
        "Citation checks support review; the team still assesses source quality and "
        "business meaning.",
        17,
        MUTED,
    )
    note(
        s,
        "Executive question: Why should we trust it? These are review controls, not a "
        "guarantee. Numerical comparisons need compatible inputs. Different definitions "
        "can explain apparent disagreements; genuine disagreements stay separate. Unknown "
        "dates remain unknown. Unsupported claims may be omitted, and responses may be partial "
        "or evidence-only. Source quotations and calculation inputs provide an audit path. "
        "The example on slide 4 shows why comparing unlike periods can mislead.",
    )

    s = base(
        prs,
        6,
        "adoption",
        "A repeatable routine, with a human at the decision",
        "Pilot one recurring meeting question and measure whether preparation becomes easier.",
    )
    for i, (title, body) in enumerate(
        [
            ("1  Refresh", "Collect current evidence and establish the comparison baseline."),
            ("2  Review", "Read the brief, challenge the key claim and discuss the gaps."),
            (
                "3  Decide",
                "Agree the team position and what to watch; carry the cited brief forward.",
            ),
        ]
    ):
        x = 0.7 + i * 4.15
        panel(s, x, 2.75, 3.9, 1.8)
        box(s, x + 0.22, 2.98, 3.45, 0.4, title, 20, TEAL, True)
        box(s, x + 0.22, 3.58, 3.43, 0.8, body, 17)
    box(s, 0.75, 4.95, 5.7, 0.35, "Available in this PoC", 19, TEAL, True)
    box(
        s,
        0.75,
        5.45,
        5.65,
        0.9,
        "On-demand research and refresh, evidence comparison, cited briefs and follow-up "
        "questions.",
        18,
    )
    box(s, 7.0, 4.95, 5.6, 0.35, "Defined scope", 19, TEAL, True)
    box(
        s,
        7.0,
        5.45,
        5.55,
        0.9,
        "Accessible sources; human review. Scheduled alerts, private-data connectors and "
        "autonomous decisions are outside scope.",
        18,
    )
    box(
        s,
        0.75,
        6.55,
        11.8,
        0.3,
        "Success measure: time to a usable brief, time to verify a claim and value to the meeting.",
        16,
        MUTED,
    )
    note(
        s,
        "Executive question: How do we adopt it? This is a working, bounded research PoC. "
        "Choose one recurring question, assign a reviewer and compare the workflow with the "
        "team's current process. Refresh is manual; first use creates a baseline. Source access "
        "and reporting dates constrain coverage. It does not provide guaranteed forecasts or "
        "unattended monitoring. Documents and briefs persist, while conversation context is "
        "in memory. Expand scope only after the pilot demonstrates a need. No measured time "
        "savings are claimed. README contains operator setup and validation instructions.",
    )
    OUT.parent.mkdir(exist_ok=True)
    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
